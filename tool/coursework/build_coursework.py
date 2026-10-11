"""Сборка курсовой работы в DOCX (и PDF) из Markdown-файлов docs/coursework.

Что делает:
1. Создаёт приложения (06_appendices.md) из актуальных исходных файлов.
2. Рисует диаграммы Mermaid в PNG (docs/coursework/diagrams/), если указан mmdc.
   Без mmdc используются уже нарисованные PNG.
3. Собирает текст в один документ через pandoc.
4. Оформляет документ (python-docx): Times New Roman 14, интервал 1,5,
   абзацный отступ 1,25 см, поля 30/15/20/20 мм, нумерация страниц,
   титульный лист, содержание, подписи рисунков и таблиц.
5. По желанию (--pdf) обновляет содержание и сохраняет PDF через LibreOffice.

Запуск из корня репозитория:
    python3 tool/coursework/build_coursework.py [--mmdc путь] [--puppeteer-config файл] [--pdf]

Нужны: pandoc, python-docx, Pillow; для --pdf — LibreOffice (модуль uno).
"""

import argparse
import os
import re
import shutil
import subprocess
import sys

from PIL import Image
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SRC = os.path.join(ROOT, "docs", "coursework")
BUILD = os.path.join(ROOT, "build", "coursework")
OUT_DOCX = os.path.join(SRC, "Курсовая_Маршрут.docx")
OUT_PDF = os.path.join(SRC, "Курсовая_Маршрут.pdf")

# Порядок разделов в итоговом документе.
PARTS = [
    "01_introduction.md",
    "02_chapter1_theory.md",
    "04_chapter2_practice.md",
    "05_conclusion.md",
    "03_references.md",
    "06_appendices.md",
]

# Приложения: буква (ГОСТ 2.105 пропускает Ё, З, Й, О, Ч, Ъ, Ы, Ь), название,
# файл и строка, с которой начинать листинг (None — с начала файла).
APPENDICES = [
    ("А", "Топологическая сортировка тем (алгоритм Кана)", "lib/domain/planner/topological_sort.dart", None),
    ("Б", "Хэширование паролей PBKDF2-HMAC-SHA256", "lib/core/utils/password_hasher.dart", None),
    ("В", "Раскладка тем по дням и интервальные повторения", "lib/domain/planner/scheduler.dart", None),
    ("Г", "Генератор плана и проверка дедлайна", "lib/domain/planner/plan_generator.dart", None),
    ("Д", "Перепланирование", "lib/domain/planner/replanner.dart", None),
    ("Е", "Таблицы базы данных", "lib/data/db/app_database.dart", None),
    ("Ж", "Правило доступа к экранам и маршруты", "lib/core/router/app_router.dart", "/// Маршруты, доступные без входа"),
    ("И", "Модульные тесты топологической сортировки", "test/domain/topological_sort_test.dart", None),
]

CHECK = "[ПРОВЕРИТЬ"


# ---------------------------------------------------------------------------
# 1. Приложения
# ---------------------------------------------------------------------------

def build_appendices():
    out = []
    for letter, title, path, start in APPENDICES:
        with open(os.path.join(ROOT, path), encoding="utf-8") as f:
            code = f.read()
        if start:
            code = code[code.index(start):]
        out.append(f"# ПРИЛОЖЕНИЕ {letter}. {title.upper()}\n")
        out.append(f"Файл `{path}`.\n")
        out.append("```dart\n" + code.rstrip() + "\n```\n")
    text = (
        "<!-- Файл создан tool/coursework/build_coursework.py из исходного кода. "
        "Не редактируйте вручную. -->\n\n" + "\n".join(out)
    )
    with open(os.path.join(SRC, "06_appendices.md"), "w", encoding="utf-8") as f:
        f.write(text)


# ---------------------------------------------------------------------------
# 2. Диаграммы Mermaid
# ---------------------------------------------------------------------------

MERMAID = re.compile(r"```mermaid\n(.*?)```", re.S)


def render_mermaid(name, text, mmdc, puppeteer):
    """Заменяет блоки mermaid ссылками на PNG и при наличии mmdc рисует их."""
    out_dir = os.path.join(SRC, "diagrams")
    os.makedirs(out_dir, exist_ok=True)
    counter = [0]

    def repl(m):
        counter[0] += 1
        base = f"{os.path.splitext(name)[0]}_{counter[0]}"
        png = os.path.join(out_dir, base + ".png")
        if mmdc:
            src = os.path.join(BUILD, base + ".mmd")
            with open(src, "w", encoding="utf-8") as f:
                f.write(m.group(1))
            cmd = [mmdc, "-i", src, "-o", png, "-s", "3", "-b", "white"]
            if puppeteer:
                cmd += ["-p", puppeteer]
            subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL)
        if not os.path.exists(png):
            sys.exit(f"Нет {png}: укажите --mmdc, чтобы нарисовать диаграммы")
        return f"![](diagrams/{base}.png)"

    return MERMAID.sub(repl, text)


# ---------------------------------------------------------------------------
# 3. Размеры рисунков и скриншоты по два в ряд
# ---------------------------------------------------------------------------

IMAGE = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)\s*$")
CAPTION = re.compile(r"^Рисунок \d+\.\d+ — ")


def image_size(rel):
    with Image.open(os.path.join(SRC, rel)) as im:
        return im.size


def is_phone_shot(rel):
    w, h = image_size(rel)
    return rel.startswith("screenshots/") and h > w


def size_attr(rel):
    w, h = image_size(rel)
    if is_phone_shot(rel):
        return "{height=11cm}"
    # диаграмма или окно компьютера: не шире 16 см и не выше 22 см
    width = min(16.0, 22.0 * w / h)
    return "{width=%.1fcm}" % width


def layout_figures(text):
    """Подряд идущие портретные скриншоты ставит по два в ряд (таблица без рамок)."""
    lines = text.split("\n")
    figures = []  # (start, end, path, caption)
    i = 0
    while i < len(lines):
        m = IMAGE.match(lines[i])
        if m and i + 2 < len(lines) and CAPTION.match(lines[i + 2]):
            figures.append((i, i + 2, m.group(1), lines[i + 2]))
            i += 3
            continue
        i += 1

    replace = {}  # start -> (end, new lines)
    k = 0
    while k < len(figures):
        s, e, path, cap = figures[k]
        nxt = figures[k + 1] if k + 1 < len(figures) else None
        if is_phone_shot(path) and nxt and nxt[0] == e + 2:
            if is_phone_shot(nxt[2]):
                row = [
                    f"| ![]({path}){{height=10cm}} | ![]({nxt[2]}){{height=10cm}} |",
                    "|:---:|:---:|",
                    f"| {cap} | {nxt[3]} |",
                ]
                replace[s] = (nxt[1], row)
                k += 2
                continue
        replace[s] = (e, [f"![]({path}){size_attr(path)}", "", cap])
        k += 1

    out = []
    i = 0
    while i < len(lines):
        if i in replace:
            end, new = replace[i]
            out.extend(new)
            i = end + 1
        else:
            out.append(lines[i])
            i += 1
    return "\n".join(out)


def strip_comments(text):
    return re.sub(r"<!--.*?-->", "", text, flags=re.S)


def assemble(mmdc, puppeteer):
    os.makedirs(BUILD, exist_ok=True)
    chunks = []
    for name in PARTS:
        with open(os.path.join(SRC, name), encoding="utf-8") as f:
            text = strip_comments(f.read())
        text = render_mermaid(name, text, mmdc, puppeteer)
        text = layout_figures(text)
        chunks.append(text.strip())
    combined = os.path.join(BUILD, "coursework.md")
    with open(combined, "w", encoding="utf-8") as f:
        f.write("\n\n".join(chunks) + "\n")
    return combined


# ---------------------------------------------------------------------------
# 4. Оформление
# ---------------------------------------------------------------------------

FONT = "Times New Roman"
MONO = "Courier New"


def set_font(style_or_run, name, size=None, bold=None):
    font = style_or_run.font
    font.name = name
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    font.color.rgb = RGBColor(0, 0, 0)
    rpr = style_or_run.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]


def para_format(pf, align=None, indent=None, spacing=1.5, before=0, after=0):
    if align is not None:
        pf.alignment = align
    if indent is not None:
        pf.first_line_indent = indent
        pf.left_indent = Cm(0)
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def style(doc, name):
    """Стиль по имени; в файлах pandoc встроенные стили названы строчными."""
    for candidate in (name, name.lower()):
        for st in doc.styles:
            if st.name == candidate:
                return st
    return None


def setup_styles(doc):
    body = [
        "Normal", "Body Text", "First Paragraph", "Compact", "Block Text",
        "Bibliography",
    ]
    for name in body:
        st = style(doc, name)
        if st is None:
            continue
        set_font(st, FONT, 14)
        para_format(st.paragraph_format, WD_ALIGN_PARAGRAPH.JUSTIFY, Cm(1.25))
    compact = style(doc, "Compact")
    if compact is not None:
        compact.paragraph_format.first_line_indent = Cm(0)

    h1 = style(doc, "Heading 1")
    if h1 is None:
        sys.exit("В документе нет стиля Heading 1")
    set_font(h1, FONT, 14, bold=True)
    h1.font.all_caps = True
    para_format(h1.paragraph_format, WD_ALIGN_PARAGRAPH.CENTER, Cm(0), before=0, after=12)
    h1.paragraph_format.page_break_before = True
    h1.paragraph_format.keep_with_next = True
    for name in ("Heading 2", "Heading 3"):
        st = style(doc, name)
        if st is None:
            continue
        set_font(st, FONT, 14, bold=True)
        st.font.italic = False
        para_format(st.paragraph_format, WD_ALIGN_PARAGRAPH.LEFT, Cm(1.25), before=12, after=6)
        st.paragraph_format.keep_with_next = True

    code = style(doc, "Source Code")
    if code is None:
        # Без подсветки синтаксиса pandoc ссылается на стиль, но не создаёт его.
        code = doc.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    set_font(code, MONO, 10)
    para_format(code.paragraph_format, WD_ALIGN_PARAGRAPH.LEFT, Cm(0), spacing=1.0, after=6)
    verb = style(doc, "Verbatim Char")
    if verb is not None:
        set_font(verb, MONO, 12)

    for name in ("Hyperlink",):
        st = style(doc, name)
        if st is not None:
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.font.underline = False


def setup_page(doc):
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.different_first_page_header_footer = True
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(footer.add_run(), "PAGE")
        set_font(footer.runs[0], FONT, 12)


def add_field(run, instr):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(sep)
    run._r.append(end)


def cell_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs


def has_image(paragraph):
    return bool(paragraph._p.xpath(".//w:drawing"))


def set_borders(table, visible):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single" if visible else "nil")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")


def format_tables(doc):
    for table in doc.tables:
        figure = any(has_image(p) for p in cell_paragraphs(table))
        # Отступ между таблицей и следующим абзацем.
        after = table._tbl.getnext()
        if after is not None and after.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            Paragraph(after, table._parent).paragraph_format.space_before = Pt(12)
        set_borders(table, not figure)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    if figure:
                        para_format(p.paragraph_format, WD_ALIGN_PARAGRAPH.CENTER, Cm(0), spacing=1.0, after=6)
                        for run in p.runs:
                            set_font(run, FONT, 12)
                            run.font.bold = False
                    else:
                        para_format(p.paragraph_format, WD_ALIGN_PARAGRAPH.LEFT, Cm(0), spacing=1.0)
                        for run in p.runs:
                            if run.style is not None and run.style.name == "Verbatim Char":
                                set_font(run, MONO, 10)
                            else:
                                set_font(run, FONT, 12)
                            if r == 0:
                                run.font.bold = True


def format_paragraphs(doc):
    for p in doc.paragraphs:
        text = p.text
        if has_image(p):
            para_format(p.paragraph_format, WD_ALIGN_PARAGRAPH.CENTER, Cm(0), spacing=1.0, before=6)
            p.paragraph_format.keep_with_next = True
        elif CAPTION.match(text):
            para_format(p.paragraph_format, WD_ALIGN_PARAGRAPH.CENTER, Cm(0), after=6)
        elif re.match(r"^Таблица \d+\.\d+ — ", text):
            para_format(p.paragraph_format, WD_ALIGN_PARAGRAPH.LEFT, Cm(0), before=6)
            p.paragraph_format.keep_with_next = True
        elif p.style.name == "Block Text":
            p.paragraph_format.left_indent = Cm(0)


def highlight_checks(doc):
    """Выделяет пометки [ПРОВЕРИТЬ …] жёлтым, чтобы их было легко найти."""
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        paragraphs.extend(cell_paragraphs(table))
    for p in paragraphs:
        if CHECK not in p.text:
            continue
        for run in list(p.runs):
            if CHECK not in run.text:
                continue
            text = run.text
            start = text.index(CHECK)
            end = text.find("]", start)
            end = len(text) if end == -1 else end + 1
            before, mark, after = text[:start], text[start:end], text[end:]
            run.text = before
            new_mark = clone_run(run, mark)
            new_mark.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if after:
                clone_run(new_mark, after).font.highlight_color = None


def clone_run(run, text):
    import copy
    new = copy.deepcopy(run._r)
    run._r.addnext(new)
    from docx.text.run import Run
    r = Run(new, run._parent)
    r.text = text
    return r


def title_page(doc):
    """Титульный лист и содержание перед первым разделом."""
    first = doc.paragraphs[0]

    def add(text, size=14, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, caps=False):
        p = first.insert_paragraph_before()
        p.style = doc.styles["Normal"]
        para_format(p.paragraph_format, align, Cm(0), before=before)
        run = p.add_run(text)
        set_font(run, FONT, size, bold=bold)
        run.font.all_caps = caps
        if CHECK in text:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return p

    add("[ПРОВЕРИТЬ: название органа управления образованием]")
    add("[ПРОВЕРИТЬ: полное название колледжа]")
    add("КУРСОВОЙ ПРОЕКТ", size=16, bold=True, before=150)
    add("по профессиональному модулю ПМ.07 «Разработка мобильных приложений»")
    add("на тему:", before=24)
    add("«Создание приложения для автоматического формирования персонального плана обучения»", bold=True)
    add("Выполнил: студент 3 курса, группа [ПРОВЕРИТЬ]", align=WD_ALIGN_PARAGRAPH.RIGHT, before=120)
    add("[ПРОВЕРИТЬ: фамилия и инициалы студента]", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add("Руководитель: [ПРОВЕРИТЬ: фамилия и инициалы]", align=WD_ALIGN_PARAGRAPH.RIGHT, before=12)
    add("Оценка: ______________", align=WD_ALIGN_PARAGRAPH.RIGHT, before=12)
    last = add("[ПРОВЕРИТЬ: город, год]", before=110)
    last.add_run().add_break(WD_BREAK.PAGE)

    p = first.insert_paragraph_before()
    para_format(p.paragraph_format, WD_ALIGN_PARAGRAPH.CENTER, Cm(0), after=12)
    run = p.add_run("СОДЕРЖАНИЕ")
    set_font(run, FONT, 14, bold=True)
    toc = first.insert_paragraph_before()
    para_format(toc.paragraph_format, WD_ALIGN_PARAGRAPH.LEFT, Cm(0))
    add_field(toc.add_run(), 'TOC \\o "1-2" \\h \\z \\u')

    # Word обновит содержание и номера страниц при открытии файла.
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def postprocess(raw, out):
    doc = Document(raw)
    setup_styles(doc)
    setup_page(doc)
    format_paragraphs(doc)
    format_tables(doc)
    highlight_checks(doc)
    title_page(doc)
    doc.save(out)


# ---------------------------------------------------------------------------
# 5. PDF через LibreOffice
# ---------------------------------------------------------------------------

UNO_SCRIPT = r'''
import sys, time, uno
from com.sun.star.beans import PropertyValue

def prop(name, value):
    p = PropertyValue(); p.Name = name; p.Value = value; return p

local = uno.getComponentContext()
resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
for _ in range(60):
    try:
        ctx = resolver.resolve("uno:socket,host=localhost,port=2002;urp;StarOffice.ComponentContext")
        break
    except Exception:
        time.sleep(0.5)
desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(sys.argv[1]), "_blank", 0, (prop("Hidden", True),))
indexes = doc.getDocumentIndexes()
if indexes.getCount() == 0:
    # Поле TOC из Word без готового текста LibreOffice не превращает в
    # оглавление, поэтому вставляем оглавление после заголовка «СОДЕРЖАНИЕ».
    text = doc.getText()
    paragraphs = text.createEnumeration()
    while paragraphs.hasMoreElements():
        par = paragraphs.nextElement()
        if par.supportsService("com.sun.star.text.Paragraph") and par.getString().strip() == "СОДЕРЖАНИЕ":
            index = doc.createInstance("com.sun.star.text.ContentIndex")
            index.CreateFromOutline = True
            index.Level = 2
            index.Title = ""
            text.insertTextContent(text.createTextCursorByRange(par.getEnd()), index, False)
            break
    indexes = doc.getDocumentIndexes()
for i in range(indexes.getCount()):
    indexes.getByIndex(i).update()
doc.refresh()
for i in range(indexes.getCount()):
    indexes.getByIndex(i).update()
doc.storeToURL(uno.systemPathToFileUrl(sys.argv[2]), (prop("FilterName", "writer_pdf_Export"),))
doc.close(True)
'''


def export_pdf(docx, pdf):
    script = os.path.join(BUILD, "export_pdf.py")
    with open(script, "w", encoding="utf-8") as f:
        f.write(UNO_SCRIPT)
    office = subprocess.Popen([
        "soffice", "--headless", "--invisible", "--norestore",
        "--accept=socket,host=localhost,port=2002;urp;",
    ])
    try:
        subprocess.run([sys.executable, script, docx, pdf], check=True, timeout=300)
    finally:
        office.terminate()
        office.wait(timeout=30)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--mmdc", help="путь к mmdc (@mermaid-js/mermaid-cli)")
    ap.add_argument("--puppeteer-config", help="настройки браузера для mmdc")
    ap.add_argument("--pdf", action="store_true", help="сохранить также PDF")
    args = ap.parse_args()

    build_appendices()
    combined = assemble(args.mmdc, args.puppeteer_config)
    raw = os.path.join(BUILD, "raw.docx")
    subprocess.run(
        ["pandoc", combined, "-f", "markdown", "-t", "docx", "--no-highlight",
         "--resource-path", SRC, "-o", raw],
        check=True,
    )
    postprocess(raw, OUT_DOCX)
    print("DOCX:", os.path.relpath(OUT_DOCX, ROOT))
    if args.pdf:
        export_pdf(OUT_DOCX, OUT_PDF)
        print("PDF:", os.path.relpath(OUT_PDF, ROOT))


if __name__ == "__main__":
    main()
